#!/usr/bin/env python3
"""Le tableau de suivi (xlsx) d'un kit, à partir d'un JSON :
{"role": "…", "lignes": [[numero, phase, test, resultat], …]}.

Usage : python3 suivi.py entree.json sortie.xlsx
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


def fabriquer(entree: str, sortie: str) -> None:
    with open(entree, encoding="utf-8") as f:
        d = json.load(f)
    wb = Workbook()
    ws = wb.active
    ws.title = d["role"][:31]
    ws["A1"] = f"Suivi des tests — {d['role']}"
    ws["A1"].font = Font(bold=True, size=14, color="172554")
    ws["A2"] = "Statut : OK · KO · Bloqué · Non fait · À refaire.  Réf. Gerimmo : les 8 caractères donnés par « Aide et retours »."
    ws["A2"].font = Font(italic=True, color="475569")
    entetes = ["N°", "Phase", "Test", "Résultat attendu", "Statut", "Réf. Gerimmo", "Date", "Commentaire"]
    ws.append([])
    ws.append(entetes)
    fin = Side(style="thin", color="CBD5E1")
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="172554")
        c.alignment = Alignment(vertical="center")
    for num, phase, test, resultat in d["lignes"]:
        ws.append([num, phase, test, resultat, None, None, None, None])
    derniere = ws.max_row
    for ligne in ws.iter_rows(min_row=5, max_row=derniere):
        for c in ligne:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=fin)
        ligne[0].font = Font(bold=True, color="1D4ED8")
    for col, largeur in zip("ABCDEFGH", (11, 22, 46, 60, 13, 14, 12, 45)):
        ws.column_dimensions[col].width = largeur
    ws.freeze_panes = "B5"
    plage = f"E5:E{derniere + 1}"
    dv = DataValidation(type="list", formula1='"OK,KO,Bloqué,Non fait,À refaire"', allow_blank=True)
    dv.add(plage)
    ws.add_data_validation(dv)
    dates = DataValidation(type="date", allow_blank=True)
    dates.add(f"G5:G{derniere + 1}")
    ws.add_data_validation(dates)
    couleurs = {"OK": "DCFCE7", "KO": "FEE2E2", "Bloqué": "FEF3C7", "À refaire": "DBEAFE"}
    for statut, fond in couleurs.items():
        ws.conditional_formatting.add(plage, CellIsRule(operator="equal", formula=[f'"{statut}"'], fill=PatternFill("solid", fgColor=fond)))
    ws.print_title_rows = "4:4"
    wb.save(sortie)


if __name__ == "__main__":
    fabriquer(sys.argv[1], sys.argv[2])
