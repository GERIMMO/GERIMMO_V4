"""Rendu des fiches de tests (PDF + Markdown) et des tableaux de suivi (Excel).

Un test est un dictionnaire :
    id, titre, phase, duree, appareil, etapes (liste), attendu,
    et en option : objectif, prerequis, attendre (dépendance à un autre
    testeur), donnees, fichiers, attention, paiement (bool), optionnel (bool).
Les textes acceptent le Markdown et les raccourcis ⟦b|…⟧ ⟦m|…⟧ ⟦f|…⟧ ⟦v|…⟧ ⟦c|…⟧.
"""

from __future__ import annotations

from pathlib import Path

from gabarits import e, md_vers_html, raccourcis_md

ICONES_APPAREIL = {"ordinateur": "💻 Ordinateur", "téléphone": "📱 Téléphone", "les deux": "💻 📱 Ordinateur ou téléphone"}


def _md_inline(texte: str) -> str:
    """Markdown d'une ligne → HTML sans le <p> englobant."""
    html = md_vers_html(texte).strip()
    if html.startswith("<p>") and html.endswith("</p>") and html.count("<p>") == 1:
        html = html[3:-4]
    return html


def carte_html(t: dict) -> str:
    pastilles = []
    if t.get("paiement"):
        pastilles.append('<span class="pastille paiement">💳 Paiement réel</span>')
    if t.get("optionnel"):
        pastilles.append('<span class="pastille claire">Facultatif</span>')
    if t.get("attendre"):
        pastilles.append('<span class="pastille alerte">⏳ Dépend d\'un autre testeur</span>')
    meta = [f"<b>{e(ICONES_APPAREIL.get(t.get('appareil', 'les deux'), t.get('appareil', '')))}</b>"]
    if t.get("prerequis"):
        meta.append(f"Avant : <b>{e(t['prerequis'])}</b>")
    blocs = []
    if t.get("objectif"):
        blocs.append(f'<p><i>{_md_inline(t["objectif"])}</i></p>')
    if t.get("attendre"):
        blocs.append(f'<div class="encadre alerte" style="margin:4px 0 6px;padding:6px 10px">⏳ {_md_inline(t["attendre"])}</div>')
    if t.get("donnees"):
        blocs.append(f'<p><b>Données :</b> {_md_inline(t["donnees"])}</p>')
    if t.get("fichiers"):
        fichiers = " ".join(f'<code class="fichier">{e(f)}</code>' for f in t["fichiers"])
        blocs.append(f"<p><b>Fichiers :</b> {fichiers}</p>")
    etapes = "".join(f"<li>{_md_inline(x)}</li>" for x in t["etapes"])
    blocs.append(f'<ol class="etapes">{etapes}</ol>')
    if t.get("attention"):
        blocs.append(f'<div class="encadre danger" style="margin:4px 0;padding:6px 10px">{_md_inline(t["attention"])}</div>')
    blocs.append(f'<div class="attendu"><b>Résultat attendu :</b> {_md_inline(t["attendu"])}</div>')
    blocs.append(
        '<div class="verdict"><span><span class="case"></span>OK</span><span><span class="case"></span>KO</span>'
        '<span><span class="case"></span>Bloqué</span><span><span class="case"></span>Non fait</span>'
        '<span>Réf. Gerimmo : <span class="ligne-ecrire"></span></span></div>'
    )
    return (
        f'<div class="test"><div class="entete"><span class="id">{e(t["id"])}</span>'
        f'<span class="nom">{e(t["titre"])}</span>{"".join(pastilles)}<span class="duree">⏱ {e(t.get("duree", ""))}</span></div>'
        f'<div class="corps"><div class="meta">{" · ".join(meta)}</div>{"".join(blocs)}</div></div>'
    )


def carte_md(t: dict) -> str:
    lignes = [f"#### {t['id']} — {t['titre']}", ""]
    infos = [ICONES_APPAREIL.get(t.get("appareil", "les deux"), t.get("appareil", "")), f"⏱ {t.get('duree', '')}"]
    if t.get("paiement"):
        infos.append("💳 **paiement réel**")
    if t.get("optionnel"):
        infos.append("facultatif")
    if t.get("prerequis"):
        infos.append(f"avant : {t['prerequis']}")
    lignes.append(" · ".join(infos))
    lignes.append("")
    if t.get("objectif"):
        lignes += [f"*{raccourcis_md(t['objectif'])}*", ""]
    if t.get("attendre"):
        lignes += [f"> ⏳ {raccourcis_md(t['attendre'])}", ""]
    if t.get("donnees"):
        lignes += [f"**Données :** {raccourcis_md(t['donnees'])}", ""]
    if t.get("fichiers"):
        lignes += ["**Fichiers :** " + ", ".join(f"`{f}`" for f in t["fichiers"]), ""]
    for i, x in enumerate(t["etapes"], 1):
        lignes.append(f"{i}. {raccourcis_md(x)}")
    lignes.append("")
    if t.get("attention"):
        lignes += [f"> [!warning]\n> {raccourcis_md(t['attention'])}", ""]
    lignes += [f"**Résultat attendu :** {raccourcis_md(t['attendu'])}", "", "- [ ] OK  - [ ] KO  - [ ] Bloqué — Réf. Gerimmo : ______", ""]
    return "\n".join(lignes)


def sommaire_html(tests: list[dict]) -> str:
    rangs = []
    phase = None
    for t in tests:
        if t["phase"] != phase:
            phase = t["phase"]
            rangs.append(f'<tr><td colspan="3" style="background:#eef2fa;font-family:Manrope;font-weight:800;color:#0f2352">{e(phase)}</td></tr>')
        marques = ("💳 " if t.get("paiement") else "") + ("⏳ " if t.get("attendre") else "") + ("(facultatif) " if t.get("optionnel") else "")
        rangs.append(f"<tr><td>{e(t['id'])}</td><td>{e(marques)}{e(t['titre'])}</td><td class='petit'>{e(t.get('duree', ''))}</td></tr>")
    return f'<table class="sommaire"><tr><th>N°</th><th>Test</th><th style="width:70px">Durée</th></tr>{"".join(rangs)}</table>'


def fiche_html(tests: list[dict]) -> str:
    morceaux = []
    phase = None
    for t in tests:
        if t["phase"] != phase:
            phase = t["phase"]
            morceaux.append(f"<h2>{e(phase)}</h2>")
        morceaux.append(carte_html(t))
    return "\n".join(morceaux)


def fiche_md(tests: list[dict]) -> str:
    morceaux = []
    phase = None
    for t in tests:
        if t["phase"] != phase:
            phase = t["phase"]
            morceaux.append(f"### {phase}\n")
        morceaux.append(carte_md(t))
    return "\n".join(morceaux)


# ------------------------------------------------------------------ Excel

def suivi_xlsx(feuilles: list[tuple[str, str, list[dict]]], destination: Path) -> Path:
    """Un classeur de suivi : une feuille par persona (nom court, titre, tests)."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    classeur = Workbook()
    classeur.remove(classeur.active)
    fin = Side(style="thin", color="D9E1EF")
    bord = Border(left=fin, right=fin, top=fin, bottom=fin)
    entete_fill = PatternFill("solid", fgColor="0F2352")
    phase_fill = PatternFill("solid", fgColor="EEF2FA")
    colonnes = [("N°", 11), ("Phase", 22), ("Test", 46), ("Résultat attendu", 60), ("Statut", 13),
                ("Réf. Gerimmo", 14), ("Date", 12), ("Commentaire", 45)]

    for nom_court, titre, tests in feuilles:
        f = classeur.create_sheet(nom_court[:31])
        f["A1"] = titre
        f["A1"].font = Font(name="Calibri", size=14, bold=True, color="0F2352")
        f["A2"] = "Statut : OK · KO · Bloqué · Non fait · À refaire.  Réf. Gerimmo : les 8 caractères donnés par « Aide et retours »."
        f["A2"].font = Font(name="Calibri", size=9, italic=True, color="4B5870")
        for i, (libelle, largeur) in enumerate(colonnes, 1):
            c = f.cell(row=4, column=i, value=libelle)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = entete_fill
            c.alignment = Alignment(vertical="center")
            c.border = bord
            f.column_dimensions[c.column_letter].width = largeur
        ligne = 5
        for t in tests:
            valeurs = [t["id"], t["phase"], ("💳 " if t.get("paiement") else "") + t["titre"] + (" (facultatif)" if t.get("optionnel") else ""),
                       _texte_brut(t["attendu"]), "", "", "", ""]
            for i, v in enumerate(valeurs, 1):
                c = f.cell(row=ligne, column=i, value=v)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = bord
                if i == 1:
                    c.font = Font(bold=True, color="2457F5")
                if i == 2:
                    c.fill = phase_fill
            ligne += 1
        validation = DataValidation(type="list", formula1='"OK,KO,Bloqué,Non fait,À refaire"', allow_blank=True)
        validation.error = "Choisissez OK, KO, Bloqué, Non fait ou À refaire."
        f.add_data_validation(validation)
        validation.add(f"E5:E{ligne}")
        date_valide = DataValidation(type="date", allow_blank=True)
        f.add_data_validation(date_valide)
        date_valide.add(f"G5:G{ligne}")
        for valeur, couleur in (("OK", "DCFCE7"), ("KO", "FEE2E2"), ("Bloqué", "FFEDD5"), ("À refaire", "FEF9C3")):
            f.conditional_formatting.add(
                f"E5:E{ligne}", CellIsRule(operator="equal", formula=[f'"{valeur}"'], fill=PatternFill("solid", fgColor=couleur))
            )
        f.freeze_panes = "B5"
        f.auto_filter.ref = f"A4:H{ligne - 1}"
        f.page_setup.orientation = "landscape"
        f.page_setup.fitToWidth = 1
        f.page_setup.fitToHeight = 0
        f.sheet_properties.pageSetUpPr.fitToPage = True
        f.print_title_rows = "4:4"
    destination.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(destination)
    return destination


def _texte_brut(texte: str) -> str:
    import re

    t = re.sub(r"⟦[bmfvc]\|(.+?)⟧", r"« \1 »", texte)
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"`(.+?)`", r"\1", t)
    return t


def affectation_xlsx(lignes: list[tuple[str, str, str]], destination: Path) -> Path:
    """Tableau d'affectation du coordinateur : qui joue quel rôle, avec quelle adresse."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    classeur = Workbook()
    f = classeur.active
    f.title = "Affectation"
    fin = Side(style="thin", color="D9E1EF")
    bord = Border(left=fin, right=fin, top=fin, bottom=fin)
    f["A1"] = "Recette Gerimmo — qui joue quel rôle"
    f["A1"].font = Font(size=14, bold=True, color="0F2352")
    f["A2"] = ("Remplir avant d'ouvrir les comptes. L'e-mail du testeur agence sert à ouvrir l'organisation ; "
               "les alias « + » servent aux personnages secondaires qu'il fait vivre.")
    f["A2"].font = Font(size=9, italic=True, color="4B5870")
    colonnes = [("Dossier", 26), ("Personnage", 26), ("Comment le compte naît", 40), ("Prénom du testeur", 18),
                ("E-mail du testeur", 32), ("Téléphone", 16), ("Appareils (PC, iPhone, Android)", 22),
                ("Kit envoyé le", 14), ("Compte actif le", 14), ("Remarques", 36)]
    for i, (libelle, largeur) in enumerate(colonnes, 1):
        c = f.cell(row=4, column=i, value=libelle)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F2352")
        c.border = bord
        c.alignment = Alignment(wrap_text=True, vertical="center")
        f.column_dimensions[c.column_letter].width = largeur
    for r, (dossier, personnage, naissance) in enumerate(lignes, 5):
        for i, v in enumerate([dossier, personnage, naissance, "", "", "", "", "", "", ""], 1):
            c = f.cell(row=r, column=i, value=v)
            c.border = bord
            c.alignment = Alignment(wrap_text=True, vertical="top")
    f.freeze_panes = "A5"
    destination.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(destination)
    return destination
